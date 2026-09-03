"""Admin-configurable .docx templates for the four outbound documents
(feasibility report, quotation, order, delivery note), each in English
and Arabic -- see backend/app/models/doc_template.py for the storage
model and backend/scripts/generate_default_doc_templates.py for how the
bundled defaults were built.

Two independent things live here:
  - Template *management*: list/upload/reset/read-back the 8
    (doc_type, language) slots -- mirrors company_logo_service.py's
    "decode-and-verify upload, uuid filename, delete old file on
    replace" shape.
  - Document *generation*: build a Jinja context from a real record
    (via the same *Out Pydantic schema its own API endpoint already
    returns) and render it through whichever template is active for
    that (doc_type, language) pair, using docxtpl.

A doc_templates row only exists for a *custom* (admin-uploaded)
template; its absence means "use the bundled default asset" -- so a
fresh install needs no seed data, same as email_template_service's
auto-created-from-default rows.
"""

import io
import uuid
from datetime import date, datetime
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.exceptions import NotFoundError, ValidationAppError
from app.models.doc_template import DocTemplate
from app.services import audit_service, pdf_generator

TABLE_NAME = "doc_templates"

DOC_TYPES = ["feasibility", "quotation", "order", "delivery_note"]
LANGUAGES = ["en", "ar"]

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

DOC_TYPE_LABELS = {
    "feasibility": "Feasibility Report",
    "quotation": "Quotation",
    "order": "Sales Order",
    "delivery_note": "Delivery Note",
}

LANGUAGE_LABELS = {"en": "English", "ar": "Arabic"}

# Documentation only, surfaced to the admin UI -- render_document() below
# never validates a template's actual content against this, same
# non-enforcing spirit as EmailTemplate's `placeholders` field. Kept in
# sync by hand with the context builders further down and with
# backend/scripts/generate_default_doc_templates.py's DEFS.
PLACEHOLDERS: dict[str, str] = {
    "feasibility": (
        "company_name, feasibility_number, customer_name, deal_number, required_by_date, checked_at, "
        "status_label, notes, generated_date, signer_name, signature_image, and a `lines` list -- each with "
        "index, product_code, product_name, quantity, supply_note, feasible_label"
    ),
    "quotation": (
        "company_name, quotation_number, customer_name, quotation_date, valid_until, status_label, "
        "subtotal_amount, discount_amount, total_amount, notes, generated_date, signer_name, signature_image, "
        "and a `lines` list -- each with index, product_code, product_name, unit, quantity, unit_price, line_total"
    ),
    "order": (
        "company_name, order_number, customer_name, order_date, requested_delivery_date, status_label, "
        "subtotal_amount, discount_amount, total_amount, notes, generated_date, signer_name, signature_image, "
        "and a `lines` list -- each with index, product_code, product_name, unit, quantity, unit_price, line_total"
    ),
    "delivery_note": (
        "company_name, delivery_note_number, order_number, customer_name, delivery_date, status_label, notes, "
        "generated_date, signer_name, signature_image, and a `lines` list -- each with index, product_code, "
        "product_name, unit, quantity_delivered"
    ),
}

# backend/app/services/doc_template_service.py -> parents[2] is backend/,
# same anchoring as company_logo_service._BACKEND_ROOT.
_BACKEND_ROOT = Path(__file__).resolve().parents[2]
_ASSETS_DIR = _BACKEND_ROOT / "app" / "assets" / "doc_templates"


def _validate_doc_type(doc_type: str) -> None:
    if doc_type not in DOC_TYPES:
        raise ValidationAppError("Unknown document type.")


def _validate_language(language: str) -> None:
    if language not in LANGUAGES:
        raise ValidationAppError("Unknown language.")


def _uploads_dir() -> Path:
    settings = get_settings()
    upload_dir = Path(settings.UPLOAD_DIR)
    if not upload_dir.is_absolute():
        upload_dir = _BACKEND_ROOT / upload_dir
    path = upload_dir / "doc_templates"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _default_asset_path(doc_type: str, language: str) -> Path:
    return _ASSETS_DIR / f"{doc_type}_{language}.docx"


def _get_row(db: Session, doc_type: str, language: str) -> DocTemplate | None:
    return (
        db.query(DocTemplate)
        .filter(DocTemplate.doc_type == doc_type, DocTemplate.language == language)
        .first()
    )


def get_active_template_path(db: Session, doc_type: str, language: str) -> Path:
    """The file docxtpl should actually load for this (doc_type,
    language): the admin-uploaded custom one if a row exists AND its
    file is still on disk, else the bundled default asset."""
    _validate_doc_type(doc_type)
    _validate_language(language)
    row = _get_row(db, doc_type, language)
    if row is not None:
        custom_path = _uploads_dir() / row.filename
        if custom_path.is_file():
            return custom_path
    default_path = _default_asset_path(doc_type, language)
    if not default_path.is_file():
        raise NotFoundError("Document template")
    return default_path


def _slot_for(db: Session, doc_type: str, language: str) -> dict:
    row = _get_row(db, doc_type, language)
    return {
        "doc_type": doc_type,
        "doc_type_label": DOC_TYPE_LABELS[doc_type],
        "language": language,
        "language_label": LANGUAGE_LABELS[language],
        "is_custom": row is not None,
        "original_filename": row.original_filename if row else None,
        "updated_at": row.updated_at if row else None,
        "placeholders": PLACEHOLDERS[doc_type],
    }


def list_templates(db: Session) -> list[dict]:
    return [_slot_for(db, doc_type, language) for doc_type in DOC_TYPES for language in LANGUAGES]


def upload_template(
    db: Session,
    doc_type: str,
    language: str,
    raw_bytes: bytes,
    original_filename: str,
    user_id: int | None = None,
) -> dict:
    _validate_doc_type(doc_type)
    _validate_language(language)

    settings = get_settings()
    max_bytes = settings.DOC_TEMPLATE_MAX_UPLOAD_MB * 1024 * 1024
    if len(raw_bytes) > max_bytes:
        raise ValidationAppError(f"Template must be under {settings.DOC_TEMPLATE_MAX_UPLOAD_MB} MB.")

    # Never trust the .docx extension -- confirm it's actually an
    # openable Word document before storing it (same "decode and
    # verify the bytes" stance company_logo_service.save_logo takes
    # with Pillow for images).
    try:
        DocxDocument(io.BytesIO(raw_bytes))
    except Exception as exc:
        raise ValidationAppError("That doesn't look like a valid .docx file.") from exc

    row = _get_row(db, doc_type, language)
    is_create = row is None
    old_filename = row.filename if row else None
    new_filename = f"{doc_type}_{language}_{uuid.uuid4().hex}.docx"
    display_name = Path(original_filename).name or "template.docx"

    # Commit the DB row before writing the file, same ordering (and
    # reasoning) as company_logo_service.save_logo: a failure here
    # leaves no orphaned file the DB doesn't know about.
    if row is None:
        row = DocTemplate(
            doc_type=doc_type, language=language, filename=new_filename, original_filename=display_name
        )
        row.created_by = user_id
        db.add(row)
    else:
        row.filename = new_filename
        row.original_filename = display_name
    row.updated_by = user_id
    db.commit()
    db.refresh(row)

    (_uploads_dir() / new_filename).write_bytes(raw_bytes)
    if old_filename:
        (_uploads_dir() / old_filename).unlink(missing_ok=True)

    if is_create:
        audit_service.log_create(db, TABLE_NAME, row.id, user_id)
    else:
        audit_service.log_update(
            db, TABLE_NAME, row.id, {"original_filename": (None, display_name)}, user_id
        )

    return _slot_for(db, doc_type, language)


def reset_template(db: Session, doc_type: str, language: str, user_id: int | None = None) -> dict:
    _validate_doc_type(doc_type)
    _validate_language(language)
    row = _get_row(db, doc_type, language)
    if row is None:
        return _slot_for(db, doc_type, language)

    old_filename = row.filename
    row_id = row.id
    db.delete(row)
    db.commit()
    audit_service.log_delete(db, TABLE_NAME, row_id, user_id)

    (_uploads_dir() / old_filename).unlink(missing_ok=True)
    return _slot_for(db, doc_type, language)


def read_active_template_bytes(db: Session, doc_type: str, language: str) -> bytes:
    return get_active_template_path(db, doc_type, language).read_bytes()


def render_document(db: Session, doc_type: str, language: str, context: dict) -> bytes:
    """Renders whichever template is active for (doc_type, language)
    with `context` (see the build_*_context functions below) and
    returns the resulting .docx bytes."""
    path = get_active_template_path(db, doc_type, language)
    tpl = DocxTemplate(str(path))

    # The signature image has to be built against this specific
    # DocxTemplate instance (docxtpl binds InlineImage to the template
    # it'll be saved into), so it can't be prepared by the context
    # builder itself -- builders instead leave a `_signature_path`
    # marker (never a real placeholder name) for this function to
    # resolve right before rendering.
    signature_path = context.pop("_signature_path", None)
    if signature_path is not None and Path(signature_path).is_file():
        context["signature_image"] = InlineImage(tpl, str(signature_path), width=Mm(35))
    else:
        context["signature_image"] = ""

    try:
        tpl.render(context)
    except Exception as exc:
        # A custom template's Jinja content is effectively untrusted
        # user input (an admin hand-edited it in Word) -- a typo like an
        # unmatched {% %} or looping over a field that isn't a list must
        # read as "fix your template", not 500 the whole request.
        raise ValidationAppError(
            "This template has an error and can't be used to generate a document. Check its "
            "placeholders, or reset it to the default under Admin -> Documents -> Document Templates."
        ) from exc

    buffer = io.BytesIO()
    tpl.save(buffer)
    return buffer.getvalue()


# --- Context builders -----------------------------------------------------
# Each builds its Jinja context from the same *Out Pydantic schema its
# own API endpoint already returns (FeasibilityOut, QuotationOut, ...)
# rather than re-deriving field names from the raw ORM a second time --
# .model_dump() gives exactly the already-shaped, display-ready fields
# those schemas expose. Company info and signer come from
# pdf_generator's own public helpers, reused as-is so a Word export and
# the existing PDF export always agree on both.


def _fmt_date(value) -> str:
    if value is None:
        return "-"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d-%b-%Y")
    return str(value)


def _fmt_money(value) -> str:
    if value is None:
        return "-"
    return f"{float(value):,.2f}"


def _fmt_qty(value) -> str:
    if value is None:
        return "-"
    return f"{float(value):g}"


def build_feasibility_context(db: Session, feasibility) -> dict:
    from app.schemas.feasibility import FeasibilityOut

    data = FeasibilityOut.from_model(feasibility).model_dump()
    company = pdf_generator.get_company_settings(db)
    signer_name, signature_path = pdf_generator.resolve_signer(db, feasibility.created_by)

    lines = []
    for idx, line in enumerate(data["lines"], start=1):
        if line["is_feasible"] is None:
            supply_note, feasible_label = "Not yet run", "-"
        elif line["bom_missing"]:
            supply_note = "No formula (BOM) set up for this product"
            feasible_label = "Not feasible -- no BOM"
        elif line["is_feasible"]:
            covered = line["covered_by_stock"] or 0
            supply_note = f"{_fmt_qty(covered)} ready in stock now" if covered else "Available to produce"
            feasible_label = "Feasible"
        else:
            shortfall_bits = [f"{s['code']} short {s['shortfall']} {s['unit']}" for s in line["shortfalls"]]
            supply_note = "; ".join(shortfall_bits) if shortfall_bits else "Short on production capacity"
            feasible_label = "Not feasible"
        lines.append(
            {
                "index": idx,
                "product_code": line["product_code"] or "",
                "product_name": line["product_name"] or f"#{line['product_id']}",
                "quantity": _fmt_qty(line["quantity"]),
                "supply_note": supply_note,
                "feasible_label": feasible_label,
            }
        )

    return {
        "company_name": company["company_name"],
        "feasibility_number": data["feasibility_number"],
        "customer_name": data["customer_name"] or "-",
        "deal_number": data["deal_number"] or "-",
        "required_by_date": _fmt_date(data["required_by_date"]),
        "checked_at": _fmt_date(data["checked_at"]),
        "status_label": data["status"].replace("_", " ").title(),
        "notes": data["notes"] or "",
        "lines": lines,
        "generated_date": _fmt_date(date.today()),
        "signer_name": signer_name or "Authorized Signatory",
        "_signature_path": signature_path,
    }


def build_quotation_context(db: Session, quotation) -> dict:
    from app.schemas.quotation import QuotationOut

    data = QuotationOut.from_model(quotation).model_dump()
    company = pdf_generator.get_company_settings(db)
    signer_name, signature_path = pdf_generator.resolve_signer(db, quotation.created_by)

    lines = [
        {
            "index": idx,
            "product_code": line["product_code"] or "",
            "product_name": line["product_name"] or f"#{line['product_id']}",
            "unit": line["unit"] or "",
            "quantity": _fmt_qty(line["quantity"]),
            "unit_price": _fmt_money(line["unit_price"]),
            "line_total": _fmt_money(line["line_total"]),
        }
        for idx, line in enumerate(data["lines"], start=1)
    ]

    return {
        "company_name": company["company_name"],
        "quotation_number": data["quotation_number"],
        "customer_name": data["customer_name"] or "-",
        "quotation_date": _fmt_date(data["quotation_date"]),
        "valid_until": _fmt_date(data["valid_until"]),
        "status_label": data["status"].replace("_", " ").title(),
        "lines": lines,
        "subtotal_amount": _fmt_money(data["subtotal_amount"]),
        "discount_amount": _fmt_money(data["discount_amount"]),
        "total_amount": _fmt_money(data["total_amount"]),
        "notes": data["notes"] or "",
        "generated_date": _fmt_date(date.today()),
        "signer_name": signer_name or "Authorized Signatory",
        "_signature_path": signature_path,
    }


def build_order_context(db: Session, order) -> dict:
    from app.schemas.order import OrderOut

    data = OrderOut.from_model(order).model_dump()
    company = pdf_generator.get_company_settings(db)
    signer_name, signature_path = pdf_generator.resolve_signer(db, order.created_by)

    lines = [
        {
            "index": idx,
            "product_code": line["product_code"] or "",
            "product_name": line["product_name"] or f"#{line['product_id']}",
            "unit": line["unit"] or "",
            "quantity": _fmt_qty(line["quantity"]),
            "unit_price": _fmt_money(line["unit_price"]),
            "line_total": _fmt_money(line["line_total"]),
        }
        for idx, line in enumerate(data["lines"], start=1)
    ]

    return {
        "company_name": company["company_name"],
        "order_number": data["order_number"],
        "customer_name": data["customer_name"] or "-",
        "order_date": _fmt_date(data["order_date"]),
        "requested_delivery_date": _fmt_date(data["requested_delivery_date"]),
        "status_label": data["status"].replace("_", " ").title(),
        "lines": lines,
        "subtotal_amount": _fmt_money(data["subtotal_amount"]),
        "discount_amount": _fmt_money(data["discount_amount"]),
        "total_amount": _fmt_money(data["total_amount"]),
        "notes": data["notes"] or "",
        "generated_date": _fmt_date(date.today()),
        "signer_name": signer_name or "Authorized Signatory",
        "_signature_path": signature_path,
    }


def build_delivery_note_context(db: Session, note) -> dict:
    from app.schemas.delivery_note import DeliveryNoteOut

    data = DeliveryNoteOut.from_model(note).model_dump()
    company = pdf_generator.get_company_settings(db)
    signer_name, signature_path = pdf_generator.resolve_signer(db, note.created_by)

    lines = [
        {
            "index": idx,
            "product_code": line["product_code"] or "",
            "product_name": line["product_name"] or f"#{line['product_id']}",
            "unit": line["unit"] or "",
            "quantity_delivered": _fmt_qty(line["quantity_delivered"]),
        }
        for idx, line in enumerate(data["lines"], start=1)
    ]

    return {
        "company_name": company["company_name"],
        "delivery_note_number": data["delivery_note_number"],
        "order_number": data["order_number"] or "-",
        "customer_name": data["customer_name"] or "-",
        "delivery_date": _fmt_date(data["delivery_date"]),
        "status_label": data["status"].replace("_", " ").title(),
        "lines": lines,
        "notes": data["notes"] or "",
        "generated_date": _fmt_date(date.today()),
        "signer_name": signer_name or "Authorized Signatory",
        "_signature_path": signature_path,
    }
